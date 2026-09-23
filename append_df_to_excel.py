import os
import pandas as pd
from openpyxl import load_workbook

def append_df_to_excel(filename, df, sheet_name='Sheet1', startrow=None,
                       truncate_sheet=False, **to_excel_kwargs):
    """
    Append a DataFrame [df] to an existing Excel file [filename]
    into the [sheet_name] Sheet. If the file does not exist, it creates one.
    """
    # Check if the Excel file exists
    if os.path.isfile(filename):
        writer = pd.ExcelWriter(filename, engine='openpyxl', mode='a')
        writer.book = load_workbook(filename)

        # Specify the starting row if not given
        if startrow is None and sheet_name in writer.book.sheetnames:
            startrow = writer.book[sheet_name].max_row

        # Truncate the sheet if specified
        if truncate_sheet and sheet_name in writer.book.sheetnames:
            idx = writer.book.sheetnames.index(sheet_name)
            writer.book.remove(writer.book.worksheets[idx])
            writer.book.create_sheet(sheet_name, idx)
        
        writer.sheets = {ws.title: ws for ws in writer.book.worksheets}
    else:
        writer = pd.ExcelWriter(filename, engine='openpyxl')

    df.to_excel(writer, sheet_name, startrow=startrow if startrow is not None else 0, **to_excel_kwargs)
    writer.save()
